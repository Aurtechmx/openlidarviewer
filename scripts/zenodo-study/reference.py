#!/usr/bin/env python3
"""
reference.py: independent reader for the Zenodo 15421291 small-file study.

Reads LAS/LAZ with laspy + lazrs (no OLV code, no PDAL) and the GroundTruth
workbook with openpyxl. Versions are pinned in requirements.txt beside this
file. Writes, for each input:

  <name>.records.bin   one 40-byte canonical record per point, the layout of
                       canonical.mjs (integer X, Y, Z and the per-point
                       attributes), so OLV's decodes can be compared to it
  <name>.json          header facts, VLR/EVLR inventory, extra dimensions,
                       attribute ranges, and the input's SHA-256

Usage (run with python -I from the study's venv):
  reference.py las <file.las|laz> <out-dir> <name>
  reference.py workbook <GroundTruth.xlsx> <spheres.csv> <out-json>
"""
import hashlib
import json
import platform
import sys

import numpy as np

RECORD = np.dtype({
    'names': ['X', 'Y', 'Z', 'intensity', 'pointSourceId', 'red8', 'green8', 'blue8',
              'classification', 'classificationFlags', 'returnNumber', 'returnCount',
              'userData', 'scanAngleRaw', 'scannerChannel', 'scanDirection',
              'edgeOfFlightLine', 'gpsTime'],
    'formats': ['<i4', '<i4', '<i4', '<u2', '<u2', 'u1', 'u1', 'u1', 'u1', 'u1', 'u1', 'u1',
                'u1', '<i2', 'u1', 'u1', 'u1', '<f8'],
    'offsets': [0, 4, 8, 12, 14, 16, 17, 18, 19, 20, 21, 22, 23, 24, 26, 27, 28, 32],
    'itemsize': 40,
})


def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for block in iter(lambda: f.read(1 << 20), b''):
            h.update(block)
    return h.hexdigest()


def versions():
    import laspy
    import lazrs
    import openpyxl
    return {
        'python': platform.python_version(),
        'laspy': laspy.__version__,
        'lazrs': getattr(lazrs, '__version__', 'unknown'),
        'numpy': np.__version__,
        'openpyxl': openpyxl.__version__,
    }


def summary(a):
    a = np.asarray(a)
    return {'dtype': str(a.dtype), 'min': a.min().item(), 'max': a.max().item(),
            'distinct': int(len(np.unique(a)))}


def read_las(path, out_dir, name):
    import laspy
    las = laspy.read(path, laz_backend=laspy.LazBackend.Lazrs) if path.lower().endswith('.laz') else laspy.read(path)
    h = las.header
    pf = h.point_format
    n = int(h.point_count)
    rec = np.zeros(n, dtype=RECORD)
    rec['X'] = np.asarray(las.X)
    rec['Y'] = np.asarray(las.Y)
    rec['Z'] = np.asarray(las.Z)
    dims = set(pf.dimension_names)
    rec['intensity'] = las.intensity
    rec['pointSourceId'] = las.point_source_id
    rec['classification'] = las.classification
    flags = (np.asarray(las.synthetic, dtype=np.uint8)
             | (np.asarray(las.key_point, dtype=np.uint8) << 1)
             | (np.asarray(las.withheld, dtype=np.uint8) << 2))
    if 'overlap' in dims:
        flags = flags | (np.asarray(las.overlap, dtype=np.uint8) << 3)
    rec['classificationFlags'] = flags
    rec['returnNumber'] = las.return_number
    rec['returnCount'] = las.number_of_returns
    rec['userData'] = las.user_data
    rec['scanAngleRaw'] = np.asarray(las.scan_angle if 'scan_angle' in dims else las.scan_angle_rank)
    if 'scanner_channel' in dims:
        rec['scannerChannel'] = las.scanner_channel
    rec['scanDirection'] = las.scan_direction_flag
    rec['edgeOfFlightLine'] = las.edge_of_flight_line
    if 'gps_time' in dims:
        rec['gpsTime'] = las.gps_time
    fields = ['X', 'Y', 'Z', 'intensity', 'pointSourceId', 'classification', 'classificationFlags',
              'returnNumber', 'returnCount', 'userData', 'scanAngleRaw', 'scanDirection', 'edgeOfFlightLine']
    if 'scanner_channel' in dims:
        fields.append('scannerChannel')
    if 'gps_time' in dims:
        fields.append('gpsTime')
    rgb = None
    if 'red' in dims:
        r16 = np.stack([np.asarray(las.red), np.asarray(las.green), np.asarray(las.blue)])
        max16 = int(r16.max())
        # OLV narrows once per file: values all <= 255 are copied, otherwise
        # the high byte is kept. The reference applies the same documented
        # rule and records how much the low byte held, so the narrowing is
        # shown to be lossless or not for this file.
        rgb8 = r16.astype(np.uint8) if max16 <= 255 else (r16 >> 8).astype(np.uint8)
        rec['red8'], rec['green8'], rec['blue8'] = rgb8
        fields += ['red8', 'green8', 'blue8']
        low = r16 & 0xFF
        rgb = {
            'max16': max16,
            'lowByteNonZeroValues': int(np.count_nonzero(low)) if max16 > 255 else 0,
            'times257Values': int(np.count_nonzero((r16 >> 8) == low)),
            'values': int(r16.size),
            'red16': summary(r16[0]), 'green16': summary(r16[1]), 'blue16': summary(r16[2]),
        }
    rec.tofile(f'{out_dir}/{name}.records.bin')

    vlrs = []
    for v in h.vlrs:
        data = getattr(v, 'record_data', None)
        if data is None:
            try:
                data = v.record_data_bytes()
            except Exception:  # noqa: BLE001 - inventory only
                data = b''
        vlrs.append({'userId': v.user_id, 'recordId': v.record_id, 'description': v.description,
                     'payloadBytes': len(data)})
    evlrs = [{'userId': v.user_id, 'recordId': v.record_id, 'description': v.description}
             for v in (h.evlrs or [])]
    wkt = None
    for v in h.vlrs:
        if v.user_id == 'LASF_Projection' and v.record_id == 2112:
            s = getattr(v, 'string', None)
            wkt = s if isinstance(s, str) else None
    text_area = None
    for v in h.vlrs:
        if v.user_id == 'LASF_Spec' and v.record_id == 3:
            text_area = bytes(v.record_data).decode('utf-8', 'replace').rstrip('\x00')
    extra = [{'name': d.name, 'dtype': str(d.dtype)} for d in pf.extra_dimensions]
    extra_stats = {d['name']: summary(las[d['name']]) for d in extra}
    norm_len = None
    if {'NormalX', 'NormalY', 'NormalZ'} <= {d['name'] for d in extra}:
        ln = np.sqrt(np.asarray(las['NormalX']) ** 2 + np.asarray(las['NormalY']) ** 2 + np.asarray(las['NormalZ']) ** 2)
        norm_len = {'min': float(ln.min()), 'max': float(ln.max())}
    standard = {d: summary(las[d]) for d in pf.standard_dimension_names
                if d not in ('X', 'Y', 'Z')}
    keys = np.stack([rec['X'], rec['Y'], rec['Z']], axis=1)
    distinct_xyz = int(len(np.unique(keys, axis=0)))
    distinct_records = int(len(np.unique(rec.view(np.void))))
    info = {
        'tool': versions(),
        'command': ' '.join(['python', '-I'] + sys.argv),
        'input': {'path': path.split('/')[-1], 'sha256': sha256(path)},
        'header': {
            'version': str(h.version), 'pointFormat': pf.id, 'pointCount': n,
            'recordLength': pf.size, 'globalEncoding': int(h.global_encoding.value),
            'scale': [float(s) for s in h.scales], 'offset': [float(o) for o in h.offsets],
            'min': [float(v) for v in h.mins], 'max': [float(v) for v in h.maxs],
            'systemIdentifier': h.system_identifier, 'generatingSoftware': h.generating_software,
            'creationDate': str(h.creation_date),
        },
        'vlrs': vlrs, 'evlrs': evlrs,
        'wkt': wkt, 'textArea': text_area,
        'extraDimensions': extra, 'extraDimensionStats': extra_stats, 'normalLength': norm_len,
        'standardDimensions': standard, 'rgb': rgb,
        'integerBounds': {'min': [int(rec[a].min()) for a in 'XYZ'], 'max': [int(rec[a].max()) for a in 'XYZ']},
        'distinctXyz': distinct_xyz, 'distinctRecords': distinct_records,
        'canonical': {'file': f'{name}.records.bin', 'count': n, 'fields': fields},
    }
    with open(f'{out_dir}/{name}.json', 'w') as f:
        json.dump(info, f, indent=2)
    print(json.dumps({'name': name, 'points': n, 'pointFormat': pf.id}))


def read_workbook(xlsx, csv_path, out_json):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        sheets[ws.title] = {'header': [c for c in rows[0] if c is not None], 'dataRows': len(rows) - 1}
    ws = wb['ReferencePoints']
    rows = [r for r in ws.iter_rows(values_only=True) if r[0] is not None][1:]
    workbook_pts = {r[0]: [float(r[1]), float(r[2]), float(r[3])] for r in rows}
    raw_decimals = max(len(str(v).split('.')[-1]) for r in rows for v in r[1:4])
    csv_pts = {}
    with open(csv_path) as f:
        next(f)
        for line in f:
            name, x, y, z = line.strip().split(',')
            csv_pts[name] = [float(x), float(y), float(z)]
    common = sorted(set(workbook_pts) & set(csv_pts))
    diffs = [abs(workbook_pts[k][a] - csv_pts[k][a]) for k in common for a in range(3)]
    trees = wb['TreesPoints']
    ids = [r[0] for r in trees.iter_rows(values_only=True) if r[0] is not None][1:]
    out = {
        'tool': versions(),
        'command': ' '.join(['python', '-I'] + sys.argv),
        'input': {'path': xlsx.split('/')[-1], 'sha256': sha256(xlsx)},
        'sheets': sheets,
        'referencePoints': {
            'count': len(workbook_pts), 'names': sorted(workbook_pts),
            'storedAs': 'text cells', 'maxDecimalsInWorkbook': raw_decimals,
        },
        'treesPoints': {'count': len(ids), 'idPrefix': sorted({str(i)[:2] for i in ids})},
        'csvComparison': {
            'csvRows': len(csv_pts), 'matchedNames': len(common),
            'onlyInWorkbook': sorted(set(workbook_pts) - set(csv_pts)),
            'onlyInCsv': sorted(set(csv_pts) - set(workbook_pts)),
            'maxAbsDifferenceMetres': max(diffs) if diffs else None,
            'note': 'The CSV rounds the workbook text to 6 decimals; a difference up to 5e-7 m is rounding.',
        },
    }
    with open(out_json, 'w') as f:
        json.dump(out, f, indent=2)
    print(json.dumps(out['csvComparison']))


if __name__ == '__main__':
    if len(sys.argv) >= 5 and sys.argv[1] == 'las':
        read_las(sys.argv[2], sys.argv[3], sys.argv[4])
    elif len(sys.argv) >= 5 and sys.argv[1] == 'workbook':
        read_workbook(sys.argv[2], sys.argv[3], sys.argv[4])
    else:
        print(__doc__, file=sys.stderr)
        sys.exit(2)
